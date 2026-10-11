"""Generate the demo submissions used to test the agent.

Two versions of a (fictional) group report for a *Travel Planning Agent*:

  travel_agent_report_v1.docx  - deliberately incomplete first draft:
        * memory type and planning strategy not explained
        * claims three tools but the log only shows one tool being called
        * mentions "Figure 1/2" but has no diagrams, no screenshots
        * Part C comparison and references missing, placeholders left in
  travel_agent_report_v2.docx  - revised version that fixes most gaps
  travel_agent_report_v2.pdf   - the same revised version as a PDF

Run:  python samples/make_samples.py
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.enum.text import WD_BREAK
from docx.shared import Inches

HERE = Path(__file__).resolve().parent
OUT = HERE / "submissions"
IMG = HERE / "assets"


# --------------------------------------------------------------------------- images
def make_images():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch

    IMG.mkdir(exist_ok=True)

    def box(ax, x, y, w, h, text, color):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02", fc=color, ec="#333", lw=1.2))
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=9, wrap=True)

    def arrow(ax, x1, y1, x2, y2):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="->", lw=1.3))

    # architecture
    fig, ax = plt.subplots(figsize=(8, 4.2))
    ax.set_xlim(0, 10); ax.set_ylim(0, 5.2); ax.axis("off")
    box(ax, 0.2, 2.1, 1.8, 1, "User goal\n'Plan 3 days in Goa\nunder Rs 15,000'", "#e8f0fe")
    box(ax, 3.2, 1.6, 2.6, 2, "Orchestrator Agent\n(Llama 3.3 70B)\nPlanner + Memory", "#fff4d6")
    box(ax, 7.2, 3.9, 2.5, 0.8, "web_search (Tavily)", "#e6f4ea")
    box(ax, 7.2, 2.4, 2.5, 0.8, "calculator", "#e6f4ea")
    box(ax, 7.2, 0.9, 2.5, 0.8, "weather_api", "#e6f4ea")
    box(ax, 3.4, 0.1, 2.2, 0.8, "Final itinerary", "#fde8e8")
    arrow(ax, 2.0, 2.6, 3.2, 2.6)
    for y in (4.3, 2.8, 1.3):
        arrow(ax, 5.8, 2.6, 7.2, y)
    arrow(ax, 4.5, 1.6, 4.5, 0.9)
    ax.set_title("Figure 1: Architecture of the Travel Planning Agent", fontsize=10)
    fig.savefig(IMG / "architecture.png", dpi=110, bbox_inches="tight"); plt.close(fig)

    # tool flow
    fig, ax = plt.subplots(figsize=(8, 1.8))
    ax.set_xlim(0, 12); ax.set_ylim(0, 2); ax.axis("off")
    steps = ["Goal", "Thought", "Tool call", "Observation", "Next action", "Final answer"]
    for i, s in enumerate(steps):
        box(ax, 0.2 + i * 2, 0.6, 1.6, 0.8, s, "#eef")
        if i:
            arrow(ax, 0.2 + i * 2 - 0.4, 1.0, 0.2 + i * 2, 1.0)
    ax.set_title("Figure 2: Tool flow (ReAct loop)", fontsize=10)
    fig.savefig(IMG / "toolflow.png", dpi=110, bbox_inches="tight"); plt.close(fig)

    # fake terminal screenshot
    fig, ax = plt.subplots(figsize=(8, 2.6))
    fig.patch.set_facecolor("#1e1e1e"); ax.set_facecolor("#1e1e1e"); ax.axis("off")
    lines = [
        "> Plan a 3-day trip to Goa under Rs 15000",
        "[Thought] I need hotel prices and weather first.",
        "[Tool Call] web_search('budget hotels north goa price per night')",
        "[Observation] 3 results: Rs 1800-2500 per night ...",
        "[Tool Call] calculator('3*2200 + 4500 + 2000')",
        "[Observation] 13100",
        "[Final Answer] Day 1: Baga beach ... total Rs 13,100",
    ]
    for i, l in enumerate(lines):
        ax.text(0.01, 0.92 - i * 0.13, l, color="#9fef00" if "Tool" in l else "#ddd", family="monospace", fontsize=8, transform=ax.transAxes)
    ax.set_title("Screenshot 1: Prompt and agent output", color="#ddd", fontsize=9)
    fig.savefig(IMG / "screenshot_run.png", dpi=110, bbox_inches="tight", facecolor=fig.get_facecolor()); plt.close(fig)


# --------------------------------------------------------------------------- content
TITLE = "Travel Planning Agent"

V1 = [
    ("title", TITLE),
    ("p", "Internal Assessment Task II - Agentic AI System using Modern LLM-based Tools"),
    ("p", "Roll No: Name"),
    ("p", "Roll No: Name"),
    ("break", None),
    ("h1", "Introduction, Abstract of problem statement"),
    ("p", "Planning a trip requires collecting information about places, costs and weather from many websites. "
          "We built an agent that takes a travel goal from the user and produces a day-wise itinerary within a budget."),
    ("h1", "AI framework(s) used"),
    ("p", "We used LangChain with the OpenRouter free API."),
    ("h1", "Tools integrated with agent"),
    ("p", "The agent uses three tools: a web search tool, a calculator and a weather API."),
    ("h1", "Part A - Agent Design (Modern View)"),
    ("h2", "Goal of the agent"),
    ("p", "The goal is to produce a complete travel itinerary that respects the user's budget and dates."),
    ("h2", "Agent role"),
    ("p", "The agent acts as a travel planner that searches, compares and plans."),
    ("h2", "Tools used"),
    ("p", "Web search is used for places and prices. The calculator and weather API are planned."),
    ("h1", "Part B - Implementation (Modern AI Tools)"),
    ("p", "The agent is implemented in Python. The architecture is shown in Figure 1 and the tool flow in Figure 2."),
    ("p", "The LLM receives the goal and decides which tool to call next."),
    ("h2", "Agent interaction log"),
    ("code", "[Thought] I should look up hotels in Goa\n[Tool Call] web_search('hotels in goa')\n[Observation] Found 5 hotels\n[Final Answer] Stay at hotel A"),
    ("h1", "Discussion, takeaway"),
    ("p", "TODO"),
    ("h1", "Conclusion"),
    ("p", "The agent can plan trips automatically."),
]

V2_EXTRA_INTRO = ("p", "Manual trip planning is slow and error-prone because budgets, prices and weather must be checked together. "
                       "Our agent turns a natural-language travel goal into a budget-checked itinerary and shows every step it took.")

V2 = [
    ("title", TITLE),
    ("p", "Internal Assessment Task II - Agentic AI System using Modern LLM-based Tools"),
    ("p", "16010123001: Aarav Shah"),
    ("p", "16010123002: Diya Mehta"),
    ("p", "16010123003: Kabir Rao"),
    ("break", None),
    ("h1", "Introduction, Abstract of problem statement"),
    ("p", "Planning a trip requires collecting information about places, costs and weather from many websites. "
          "We built an agent that takes a travel goal from the user and produces a day-wise itinerary within a budget."),
    V2_EXTRA_INTRO,
    ("h1", "AI framework(s) used"),
    ("p", "We used a ReAct-style agent built with LangChain, with Llama 3.3 70B accessed through the OpenRouter free API. "
          "LangChain provides the tool-calling loop and OpenRouter gives free access to an open model."),
    ("h1", "Tools integrated with agent"),
    ("p", "The agent uses three tools: web_search (Tavily search API) for places and prices, calculator for budget arithmetic, "
          "and weather_api (Open-Meteo) for the forecast. Each tool was chosen because it supplies information the LLM cannot know reliably on its own."),
    ("h1", "Part A - Agent Design (Modern View)"),
    ("h2", "Goal of the agent"),
    ("p", "The goal is to produce a complete travel itinerary that respects the user's budget and dates."),
    ("h2", "Agent role"),
    ("p", "The agent acts as a travel planner that searches, compares, calculates and plans."),
    ("h2", "Tools used"),
    ("p", "web_search, calculator and weather_api as described above."),
    ("h2", "Memory type"),
    ("p", "Short-term conversational memory (the ReAct scratchpad of thoughts and observations) is kept in the context window; "
          "user preferences such as budget are stored in a JSON file between sessions."),
    ("h2", "Planning strategy"),
    ("p", "The agent follows the ReAct strategy: it reasons about the goal, calls a tool, observes the result and decides the next step until the itinerary is complete."),
    ("h1", "Part B - Implementation (Modern AI Tools)"),
    ("h2", "LLM-based agent"),
    ("p", "Llama 3.3 70B was chosen because it is free on OpenRouter, supports long context and follows tool-calling instructions well."),
    ("h2", "Architecture diagram"),
    ("img", "architecture.png"),
    ("caption", "Figure 1: Architecture of the Travel Planning Agent"),
    ("h2", "Tool flow diagram"),
    ("img", "toolflow.png"),
    ("caption", "Figure 2: Tool flow (ReAct loop)"),
    ("h2", "Agent interaction log"),
    ("code", "[Thought] I need hotel prices and weather first.\n[Tool Call] web_search('budget hotels north goa price per night')\n"
             "[Observation] 3 results: Rs 1800-2500 per night\n[Tool Call] weather_api('Goa', '2026-12-10')\n[Observation] Sunny, 31C\n"
             "[Tool Call] calculator('3*2200 + 4500 + 2000')\n[Observation] 13100\n[Final Answer] Day 1: Baga beach ... total Rs 13,100"),
    ("h2", "Prompts and output screenshots"),
    ("img", "screenshot_run.png"),
    ("caption", "Screenshot 1: Prompt and agent output"),
    ("h2", "Code repository"),
    ("p", "Code and notebook: https://github.com/example-group/travel-planning-agent"),
    ("h1", "Part C - Analysis & Comparison"),
    ("table", [["Classical Agent (Syllabus)", "Modern Agent (This Activity)"], ["Rule-based", "LLM-driven"],
               ["Static logic", "Dynamic planning"], ["No memory", "Contextual memory"], ["Single-step", "Multi-step reasoning"]]),
    ("p", "A rule-based planner can only follow fixed if-then rules for known cities, while our agent plans dynamically and adapts when a tool returns unexpected results."),
    ("h1", "Discussion, takeaway"),
    ("p", "Tool use made the itinerary costs reliable; the main limitation is the rate limit of free APIs and occasional malformed tool calls, which we handle with retries."),
    ("h1", "Conclusion"),
    ("p", "The agent plans trips end to end with observable reasoning and verified budget calculations."),
    ("h1", "References"),
    ("p", "[1] S. Yao et al., ReAct: Synergizing Reasoning and Acting in Language Models, ICLR 2023."),
    ("p", "[2] LangChain documentation, https://python.langchain.com"),
]


def write_docx(content, path: Path):
    d = docx.Document()
    for kind, val in content:
        if kind == "title":
            d.add_heading(val, 0)
        elif kind == "h1":
            d.add_heading(val, 1)
        elif kind == "h2":
            d.add_heading(val, 2)
        elif kind == "p":
            d.add_paragraph(val)
        elif kind == "caption":
            d.add_paragraph(val, style="Caption")
        elif kind == "code":
            for line in val.splitlines():
                p = d.add_paragraph()
                r = p.add_run(line)
                r.font.name = "Consolas"
        elif kind == "img":
            d.add_picture(str(IMG / val), width=Inches(5.5))
        elif kind == "table":
            t = d.add_table(rows=0, cols=len(val[0]))
            t.style = "Table Grid"
            for row in val:
                cells = t.add_row().cells
                for c, v in zip(cells, row):
                    c.text = v
        elif kind == "break":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.save(path)


def write_pdf(content, path: Path):
    import html

    import pymupdf

    parts = []
    for kind, val in content:
        if kind == "title":
            parts.append(f"<h1 style='text-align:center'>{html.escape(val)}</h1>")
        elif kind == "h1":
            parts.append(f"<h2>{html.escape(val)}</h2>")
        elif kind == "h2":
            parts.append(f"<h3>{html.escape(val)}</h3>")
        elif kind in ("p", "caption"):
            style = " style='font-style:italic'" if kind == "caption" else ""
            parts.append(f"<p{style}>{html.escape(val)}</p>")
        elif kind == "code":
            parts.append("<pre style='font-size:8pt'>" + html.escape(val) + "</pre>")
        elif kind == "img":
            parts.append(f"<img src='{val}' width='420'/>")
        elif kind == "table":
            rows = "".join("<tr>" + "".join(f"<td style='border:1px solid #444;padding:3px'>{html.escape(c)}</td>" for c in r) + "</tr>" for r in val)
            parts.append(f"<table style='border-collapse:collapse'>{rows}</table>")
        elif kind == "break":
            parts.append("<p style='page-break-after:always'></p>")
    story = pymupdf.Story("<body style='font-family:sans-serif;font-size:11pt'>" + "".join(parts) + "</body>",
                          archive=pymupdf.Archive(str(IMG)))
    writer = pymupdf.DocumentWriter(str(path))
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (54, 54, -54, -54)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()


def main():
    OUT.mkdir(exist_ok=True)
    make_images()
    write_docx(V1, OUT / "travel_agent_report_v1.docx")
    write_docx(V2, OUT / "travel_agent_report_v2.docx")
    write_pdf(V2, OUT / "travel_agent_report_v2.pdf")
    print("written to", OUT)


if __name__ == "__main__":
    main()
