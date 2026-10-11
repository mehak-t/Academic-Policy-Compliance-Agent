"""Document ingestion layer.

Turns PDF / DOCX / Markdown / text / Jupyter notebook files into a flat list of
blocks (headings, paragraphs, tables, code, images, captions) with page numbers,
so every later tool can point back to *where* a piece of evidence came from.
"""

from __future__ import annotations

import hashlib
import json
import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path

WORDS_PER_PAGE = 450  # used only when a format has no real pagination
CAPTION_RE = re.compile(r"^\s*(figure|fig\.|table|screenshot|diagram)\s*[\d.:\-]", re.I)

SUPPORTED = {".pdf", ".docx", ".md", ".txt", ".ipynb"}


@dataclass
class Block:
    page: int
    kind: str  # heading | para | table | code | image | caption
    text: str
    level: int = 0

    def to_dict(self):
        return {"page": self.page, "kind": self.kind, "text": self.text, "level": self.level}


@dataclass
class Document:
    doc_id: str
    path: str
    file_type: str
    pages: int
    pages_estimated: bool
    blocks: list[Block] = field(default_factory=list)
    sha1: str = ""

    # ---- convenience views -------------------------------------------------
    @property
    def headings(self) -> list[Block]:
        return [b for b in self.blocks if b.kind == "heading"]

    @property
    def image_count(self) -> int:
        return sum(1 for b in self.blocks if b.kind == "image")

    @property
    def table_count(self) -> int:
        return sum(1 for b in self.blocks if b.kind == "table")

    @property
    def code_count(self) -> int:
        return sum(1 for b in self.blocks if b.kind == "code")

    @property
    def word_count(self) -> int:
        return sum(len(b.text.split()) for b in self.blocks if b.kind != "image")

    def text(self, start_page: int = 1, end_page: int | None = None) -> str:
        end_page = end_page or self.pages
        out, last = [], None
        for b in self.blocks:
            if not (start_page <= b.page <= end_page):
                continue
            if b.page != last:
                out.append(f"\n--- page {b.page} ---")
                last = b.page
            prefix = {"heading": "#" * max(b.level, 1) + " ", "code": "```\n", "table": "[TABLE] "}.get(b.kind, "")
            suffix = "\n```" if b.kind == "code" else ""
            out.append(prefix + b.text + suffix)
        return "\n".join(out).strip()

    def summary(self) -> dict:
        return {
            "doc_id": self.doc_id,
            "file": Path(self.path).name,
            "file_type": self.file_type,
            "pages": self.pages,
            "pages_estimated": self.pages_estimated,
            "word_count": self.word_count,
            "headings": [f"p{h.page}: {h.text}" for h in self.headings][:40],
            "images": self.image_count,
            "tables": self.table_count,
            "code_blocks": self.code_count,
        }


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------

def _clean(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def _kind_for_text(text: str) -> str:
    return "caption" if CAPTION_RE.match(text) else "para"


def _read_pdf(path: Path) -> tuple[list[Block], int, bool]:
    import pymupdf

    if hasattr(pymupdf, "no_recommend_layout"):
        pymupdf.no_recommend_layout()
    blocks: list[Block] = []
    with pymupdf.open(path) as pdf:
        sizes = []
        for page in pdf:
            for b in page.get_text("dict")["blocks"]:
                for line in b.get("lines", []):
                    for span in line["spans"]:
                        if span["text"].strip():
                            sizes.append(round(span["size"], 1))
        body = statistics.median(sizes) if sizes else 11

        for pno, page in enumerate(pdf, start=1):
            for b in page.get_text("dict")["blocks"]:
                if b.get("type") == 1:
                    blocks.append(Block(pno, "image", "[image]"))
                    continue
                lines = b.get("lines", [])
                text = _clean(" ".join(s["text"] for l in lines for s in l["spans"]))
                if not text:
                    continue
                spans = [s for l in lines for s in l["spans"] if s["text"].strip()]
                max_size = max(s["size"] for s in spans)
                bold = all("bold" in s["font"].lower() or s["flags"] & 16 for s in spans)
                short = len(text.split()) <= 14
                if short and (max_size >= body * 1.18 or (bold and max_size >= body)):
                    level = 1 if max_size >= body * 1.5 else 2 if max_size >= body * 1.18 else 3
                    blocks.append(Block(pno, "heading", text, level))
                else:
                    blocks.append(Block(pno, _kind_for_text(text), text))
            # images drawn as XObjects are not always returned as image blocks
            n_xobj = len(page.get_images(full=True))
            n_seen = sum(1 for x in blocks if x.page == pno and x.kind == "image")
            for _ in range(max(0, n_xobj - n_seen)):
                blocks.append(Block(pno, "image", "[image]"))
            try:
                for t in page.find_tables().tables:
                    rows = t.extract()
                    blocks.append(Block(pno, "table", " | ".join(_clean(" ; ".join(str(c or "") for c in r)) for r in rows)[:2000]))
            except Exception:
                pass
        return blocks, pdf.page_count, False


def _read_docx(path: Path) -> tuple[list[Block], int, bool]:
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(str(path))
    blocks: list[Block] = []
    body = d.element.body
    # Pages rendered by Word are reliable; files never opened in Word only have
    # manual page breaks, so we also estimate overflow from the amount of content.
    rendered = body.find(".//" + qn("w:lastRenderedPageBreak")) is not None
    page, words_on_page = 1, 0
    mono = {"consolas", "courier new", "courier", "lucida console", "source code pro", "cascadia code"}

    def bump(words: int):
        nonlocal page, words_on_page
        words_on_page += words
        if not rendered and words_on_page > WORDS_PER_PAGE:
            page += 1
            words_on_page = 0

    for el in body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            if rendered:
                breaks = len(el.findall(".//" + qn("w:lastRenderedPageBreak")))
            else:
                breaks = sum(1 for br in el.findall(".//" + qn("w:br")) if br.get(qn("w:type")) == "page")
            if breaks:
                page += breaks
                words_on_page = 0
            para = docx.text.paragraph.Paragraph(el, d)
            text = _clean(para.text)
            n_img = len(el.findall(".//" + qn("w:drawing"))) + len(el.findall(".//" + qn("w:pict")))
            for _ in range(n_img):
                blocks.append(Block(page, "image", "[image]"))
                bump(150)  # an image takes up real space on the page
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            runs = [r for r in para.runs if r.text.strip()]
            all_bold = bool(runs) and all(r.bold for r in runs)
            is_code = "code" in style or "source" in style or (bool(runs) and all((r.font.name or "").lower() in mono for r in runs))
            if style.startswith("heading") or style == "title":
                m = re.search(r"(\d)", style)
                blocks.append(Block(page, "heading", text, int(m.group(1)) if m else 1))
            elif is_code:
                if blocks and blocks[-1].kind == "code" and blocks[-1].page == page:
                    blocks[-1].text += "\n" + para.text  # merge consecutive code lines
                else:
                    blocks.append(Block(page, "code", para.text))
            elif all_bold and len(text.split()) <= 12:
                blocks.append(Block(page, "heading", text, 3))
            else:
                blocks.append(Block(page, _kind_for_text(text), text))
            bump(len(text.split()))
        elif tag == "tbl":
            table = docx.table.Table(el, d)
            rows = []
            for r in table.rows:
                cells = []
                for c in r.cells:
                    t = _clean(c.text)
                    if not cells or cells[-1] != t:  # merged cells repeat text
                        cells.append(t)
                rows.append(" ; ".join(cells))
            text = " | ".join(rows)
            blocks.append(Block(page, "table", text[:2000]))
            bump(len(text.split()))

    return blocks, max(page, 1), not rendered


def _read_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    in_code, code_buf = False, []
    para_buf: list[str] = []

    def flush_para():
        if para_buf:
            t = _clean(" ".join(para_buf))
            if t.startswith("|"):
                blocks.append(Block(0, "table", t))
            else:
                blocks.append(Block(0, _kind_for_text(t), t))
            para_buf.clear()

    for line in text.splitlines():
        if line.strip().startswith("```"):
            if in_code:
                blocks.append(Block(0, "code", "\n".join(code_buf)))
                code_buf, in_code = [], False
            else:
                flush_para()
                in_code = True
            continue
        if in_code:
            code_buf.append(line)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            flush_para()
            blocks.append(Block(0, "heading", _clean(m.group(2)), len(m.group(1))))
        elif re.match(r"^\s*!\[.*\]\(.*\)", line):
            flush_para()
            blocks.append(Block(0, "image", "[image] " + _clean(line)))
        elif not line.strip():
            flush_para()
        elif line.strip().startswith("|") and para_buf and not para_buf[-1].strip().startswith("|"):
            flush_para()
            para_buf.append(line)
        else:
            para_buf.append(line)
    flush_para()
    if code_buf:
        blocks.append(Block(0, "code", "\n".join(code_buf)))
    return blocks


def _paginate(blocks: list[Block]) -> int:
    page, words = 1, 0
    for b in blocks:
        b.page = page
        words += 150 if b.kind == "image" else len(b.text.split())
        if words > WORDS_PER_PAGE:
            page += 1
            words = 0
    return max(page, 1)


def _read_ipynb(path: Path) -> list[Block]:
    nb = json.loads(path.read_text(encoding="utf-8"))
    blocks: list[Block] = []
    for cell in nb.get("cells", []):
        src = "".join(cell.get("source", []))
        if cell.get("cell_type") == "markdown":
            blocks.extend(_read_markdown(src))
        elif cell.get("cell_type") == "code":
            if src.strip():
                blocks.append(Block(0, "code", src))
            for out in cell.get("outputs", []):
                data = out.get("data", {})
                if "image/png" in data or "image/jpeg" in data:
                    blocks.append(Block(0, "image", "[image] notebook output"))
                text = "".join(out.get("text", []) or data.get("text/plain", []))
                if text.strip():
                    blocks.append(Block(0, "code", text[:1500]))
    return blocks


def read_document(path: str | Path, doc_id: str | None = None) -> Document:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type {ext}; supported: {sorted(SUPPORTED)}")

    if ext == ".pdf":
        blocks, pages, estimated = _read_pdf(path)
    elif ext == ".docx":
        blocks, pages, estimated = _read_docx(path)
    elif ext == ".ipynb":
        blocks = _read_ipynb(path)
        pages, estimated = _paginate(blocks), True
    else:
        blocks = _read_markdown(path.read_text(encoding="utf-8", errors="replace"))
        pages, estimated = _paginate(blocks), True

    sha1 = hashlib.sha1(path.read_bytes()).hexdigest()
    return Document(
        doc_id=doc_id or path.stem,
        path=str(path),
        file_type=ext.lstrip("."),
        pages=pages,
        pages_estimated=estimated,
        blocks=blocks,
        sha1=sha1,
    )
