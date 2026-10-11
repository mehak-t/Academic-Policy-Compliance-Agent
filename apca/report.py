"""Audit report generation (Markdown, DOCX, JSON, CSV)."""

from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

STATUS_ICON = {
    "Satisfied": "OK", "Partially Satisfied": "PARTIAL", "Not Satisfied": "MISSING",
    "Not Verifiable": "UNVERIFIED", "Not Applicable": "N/A",
}


def checklist_rows(result: dict) -> list[dict]:
    rows = []
    for r in result.get("checklist", []):
        a = result.get("assessments", {}).get(r["id"], {})
        rows.append({
            "id": r["id"],
            "requirement": r["description"],
            "rubric_category": r.get("rubric_category", ""),
            "priority": r.get("priority", ""),
            "weight": r.get("weight", 1),
            "status": a.get("status", "Not Verifiable"),
            "page": a.get("page") or "",
            "evidence": a.get("evidence", ""),
            "confidence": a.get("confidence", ""),
            "reason": a.get("reason", ""),
            "suggested_fix": a.get("suggested_fix", ""),
        })
    return rows


def _esc(s) -> str:
    return str(s).replace("|", "/").replace("\n", " ")


def to_markdown(result: dict, approved_by: str | None = None) -> str:
    sc = result.get("score") or {}
    sub = result.get("submission", {})
    pol = result.get("policy", {})
    L = [
        "# Pre-submission Compliance Audit",
        "",
        f"- **Submission:** {sub.get('file')}  ",
        f"- **Policy:** {pol.get('file')}  ",
        f"- **Run:** {result.get('run_id')} · {result.get('steps')} agent steps · model `{result.get('llm_usage', {}).get('model')}`  ",
        f"- **Generated:** {datetime.now():%d %b %Y %H:%M}" + (f" · approved by {approved_by}" if approved_by else " · NOT YET APPROVED"),
        "",
        "> Advisory audit only. It is not an official grade; please review uncertain findings.",
        "",
        "## Summary",
        "",
        str(result.get("final_answer", "")).strip(),
        "",
        f"**Compliance score:** {sc.get('score_percent', 'n/a')}%  ",
    ]
    if sc.get("rubric_estimate"):
        L.append(f"**Estimated rubric coverage:** {sc.get('estimated_total')} / {sc.get('rubric_total')}  ")
        L += ["", "| Rubric component | Max | Estimated |", "|---|---|---|"]
        for cat, m in sc["rubric_estimate"].items():
            L.append(f"| {_esc(cat)} | {m['max']:g} | {m['estimated'] if m['estimated'] is not None else 'no items'} |")
    counts = sc.get("status_counts", {})
    if counts:
        L += ["", "Status counts: " + ", ".join(f"{k}: {v}" for k, v in counts.items())]

    L += ["", "## Prioritised correction plan", ""]
    acts = result.get("actions", [])
    if not acts:
        L.append("No open items - everything required was found.")
    for a in acts:
        L.append(f"{a['rank']}. **[{a['priority']}] {a['id']} - {a['status']}:** {a['action']}")

    L += ["", "## Requirement checklist", "", "| ID | Requirement | Priority | Status | Page | Evidence / gap | Conf. |", "|---|---|---|---|---|---|---|"]
    for r in checklist_rows(result):
        gap = r["evidence"] or r["reason"]
        L.append(f"| {r['id']} | {_esc(r['requirement'])} | {r['priority']} | {STATUS_ICON.get(r['status'], r['status'])} {r['status']} | {r['page']} | {_esc(gap)[:160]} | {r['confidence']} |")

    fmt = result.get("format")
    if fmt:
        L += ["", "## Format and structure checks", "", "| Check | Expected | Actual | Result |", "|---|---|---|---|"]
        for c in fmt["checks"]:
            L.append(f"| {_esc(c['check'])} | {_esc(c['expected'])} | {_esc(c['actual'])}{_esc(c.get('note', ''))} | {'pass' if c['passed'] else 'FAIL'} |")

    con = result.get("consistency")
    if con:
        L += ["", "## Consistency findings", ""]
        if not con["findings"]:
            L.append("No contradictions or placeholders detected.")
        for f in con["findings"]:
            L.append(f"- **{f['severity']}** ({f['type']}, p{f['page'] or '?'}): {f['message']}")

    cmp_ = result.get("comparison")
    if cmp_ and cmp_.get("found"):
        L += ["", "## Comparison with previous version", "",
              f"Previous audit {cmp_['previous_audit']} ({cmp_['previous_file']}): {cmp_['previous_score']}% -> now {cmp_['current_score']}%", ""]
        for key, label in (("fixed", "Fixed"), ("improved", "Improved"), ("regressed", "Regressed"), ("still_open", "Still open")):
            items = cmp_.get(key, [])
            if items:
                L.append(f"- **{label}:** " + ", ".join(f"{i['id']} ({i['before']} -> {i['after']})" for i in items))

    if sc.get("formula"):
        L += ["", "## Scoring assumptions", "", f"`{sc['formula']}` with " + ", ".join(f"{k}={v}" for k, v in sc["status_values"].items() if v is not None) + "; Not Applicable excluded."]

    u = result.get("llm_usage", {})
    L += ["", "## Agent trace", "", f"Tools used: " + ", ".join(f"{k} x{v}" for k, v in result.get("tool_usage", {}).items()),
          f"LLM calls: {u.get('llm_calls')} · tokens in/out: {u.get('prompt_tokens')}/{u.get('completion_tokens')}",
          f"Full interaction log: `{Path(result.get('logs', {}).get('markdown', '')).name}`"]
    return "\n".join(L) + "\n"


def to_docx(result: dict, path: Path, approved_by: str | None = None):
    import docx
    from docx.shared import Pt, RGBColor

    colors = {"Satisfied": "1E7B34", "Partially Satisfied": "B7791F", "Not Satisfied": "C53030", "Not Verifiable": "4A5568", "Not Applicable": "718096"}
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading("Pre-submission Compliance Audit", 0)
    sc = result.get("score") or {}
    p = d.add_paragraph()
    p.add_run(f"Submission: {result['submission'].get('file')}\nPolicy: {result['policy'].get('file')}\n")
    p.add_run(f"Score: {sc.get('score_percent', 'n/a')}%")
    if sc.get("rubric_estimate"):
        p.add_run(f"   ·   Estimated rubric coverage: {sc['estimated_total']} / {sc['rubric_total']}")
    p.add_run("\nApproved by: " + (approved_by or "not approved")).italic = True
    d.add_paragraph("Advisory audit only - not an official grade.").italic = True

    d.add_heading("Summary", 1)
    d.add_paragraph(str(result.get("final_answer", "")))

    d.add_heading("Prioritised correction plan", 1)
    for a in result.get("actions", []):
        d.add_paragraph(f"[{a['priority']}] {a['id']} ({a['status']}): {a['action']}", style="List Number")

    d.add_heading("Requirement checklist", 1)
    rows = checklist_rows(result)
    t = d.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(["ID", "Requirement", "Status", "Evidence / gap", "Fix"]):
        t.rows[0].cells[i].text = h
    for r in rows:
        c = t.add_row().cells
        c[0].text = r["id"]
        c[1].text = r["requirement"]
        run = c[2].paragraphs[0].add_run(r["status"])
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(colors.get(r["status"], "000000"))
        c[3].text = (f"p{r['page']}: " if r["page"] else "") + (r["evidence"] or r["reason"])
        c[4].text = r["suggested_fix"]

    con = result.get("consistency") or {}
    if con.get("findings"):
        d.add_heading("Consistency findings", 1)
        for f in con["findings"]:
            d.add_paragraph(f"{f['severity']}: {f['message']}", style="List Bullet")
    d.save(path)


def export(result: dict, out_dir: str | Path, approved_by: str | None = None) -> dict:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    base = out / f"audit_{result['run_id']}"
    paths = {
        "markdown": base.with_suffix(".md"),
        "json": base.with_suffix(".json"),
        "csv": base.with_suffix(".csv"),
        "docx": base.with_suffix(".docx"),
    }
    paths["markdown"].write_text(to_markdown(result, approved_by), encoding="utf-8")
    paths["json"].write_text(json.dumps(result | {"approved_by": approved_by}, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    rows = checklist_rows(result)
    with paths["csv"].open("w", newline="", encoding="utf-8") as fh:
        if rows:
            w = csv.DictWriter(fh, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
    to_docx(result, paths["docx"], approved_by)
    return {k: str(v) for k, v in paths.items()}
